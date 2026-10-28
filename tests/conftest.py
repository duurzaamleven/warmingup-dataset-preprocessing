# SPDX-License-Identifier: EUPL-1.2
"""Gedeelde testfixtures: een kleine, zelfgebouwde zip met dezelfde vorm als
``data-en-resultaten.zip``, maar dan met gewone Deflate voor alle leden (ook ``thermostat-temp.csv``:
de Deflate64-specifieke lezer heeft zijn eigen test tegen de echte data, zie
``test_deflate64_echte_data.py``)."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

BOILER_TEMP_KOP = "House,Time,Boiler Water Supply Temperature (°C),Boiler Water Return Temperature (°C),Flow (l/h)"
THERMOSTAT_TEMP_KOP = "House,Time,Thermostat Desired Temperature (°C),Thermostat Measured Temperature (°C)"
BOILER_24H_KOP = (
    "House,Day (Europe/Amsterdam),Gas Usage Domestic Hot Water (DHW) (m³),"
    "Gas Usage Central Heating (CH) (m³),Heater Active Duration DWH (hours),"
    "Heater Active Duration CH (hours),Water Pressure CH Min (Bar -1=unknown),"
    "Water Pressure CH Avg (Bar -1=unknown),Water Pressure CH Max (Bar -1=unknown)"
)


def _boiler_temp_regels(home: str, n: int, start_uur: str = "2021-01-01T00:00:00Z") -> list[str]:
    tijden = pd.date_range(start_uur, periods=n, freq="min", tz="UTC")
    regels = []
    for i, t in enumerate(tijden):
        temp_aanvoer = 40.0 + (i % 10)
        temp_retour = 30.0 + (i % 5)
        debiet = 0 if i % 4 else 300
        regels.append(f"{home},{t.strftime('%Y-%m-%dT%H:%M:%SZ')},{temp_aanvoer:.2f},{temp_retour:.2f},{debiet}")
    return regels


def _thermostat_temp_regels(home: str, n: int, start_uur: str = "2021-01-01T00:00:00Z") -> list[str]:
    tijden = pd.date_range(start_uur, periods=n, freq="7min", tz="UTC")
    regels = []
    for i, t in enumerate(tijden):
        regels.append(f"{home},{t.strftime('%Y-%m-%dT%H:%M:%SZ')},{19.0 + (i % 3)},{19.5 + (i % 4) * 0.1}")
    return regels


def _etmaal_regels(home: str, n: int, start_dag: str = "2021-01-01") -> list[str]:
    dagen = pd.date_range(start_dag, periods=n, freq="D")
    regels = []
    for i, d in enumerate(dagen):
        regels.append(f"{home},{d.strftime('%Y-%m-%d')},0.5,3.0,0.2,1.0,1.5,1.6,1.7")
    return regels


def bouw_test_zip(
    pad: Path,
    homes: list[str],
    rijen_per_woning: int = 20,
    etmaal_dagen_per_woning: int = 5,
    ongesorteerd: bool = False,
) -> Path:
    """Bouw een kleine zip met dezelfde vijf leden als het echte archief (versie van 27-10-2026), gevuld met synthetische
    data voor ``homes``. Bij ``ongesorteerd=True`` worden de regels van de grote bestanden zo
    verdeeld dat woning A en woning B elkaar afwisselen (bewust GEEN aaneengesloten blokken), zodat
    de ongesorteerde route getest kan worden."""
    boiler_regels = []
    thermo_regels = []
    for home in homes:
        boiler_regels.append(_boiler_temp_regels(home, rijen_per_woning))
        thermo_regels.append(_thermostat_temp_regels(home, rijen_per_woning))

    def _vlecht(delen: list[list[str]]) -> list[str]:
        if not ongesorteerd or len(delen) < 2:
            return [regel for deel in delen for regel in deel]
        # Verweef in kleine brokken van 3 regels per woning, zodat elke woning meerdere keren
        # terugkomt in het bestand (expliciet NIET aaneengesloten).
        resultaat = []
        indices = [0] * len(delen)
        actief = True
        while actief:
            actief = False
            for d_idx, deel in enumerate(delen):
                brok = deel[indices[d_idx] : indices[d_idx] + 3]
                if brok:
                    resultaat.extend(brok)
                    indices[d_idx] += 3
                    actief = True
        return resultaat

    boiler_inhoud = BOILER_TEMP_KOP + "\n" + "\n".join(_vlecht(boiler_regels)) + "\n"
    thermo_inhoud = THERMOSTAT_TEMP_KOP + "\n" + "\n".join(_vlecht(thermo_regels)) + "\n"

    etmaal_regels = []
    for home in homes:
        etmaal_regels.append(_etmaal_regels(home, etmaal_dagen_per_woning))
    etmaal_inhoud = BOILER_24H_KOP + "\n" + "\n".join(_vlecht(etmaal_regels)) + "\n"

    readme_inhoud = "Boiler temp: test.\nBoiler 24h: test.\nThermostat Temp: test.\n"

    # De resultatentabel: 'info'-blad plus 'Sheet1' met de kop op rij 5 (0-geïndexeerd).
    wb = openpyxl.Workbook()
    info_ws = wb.active
    info_ws.title = "info"
    info_ws.cell(row=2, column=3, value="Testmetadata, regel 1")
    info_ws.cell(row=13, column=3, value="Testmetadata, regel 2")

    sheet1 = wb.create_sheet("Sheet1")
    kop = [
        "Order", "woning_id", "model_training_rms", "model_validation_rms",
        "model_training_mean_bias_error", "model_validation_mean_bias_error", "steady_state_heat_supply, Q_d [W]",
        "C_i [J/C]", "R_ie [C/W]", "A_w [m2]", "C_e [J/C]", "R_ea [C/W]", "sv_supply_wa_60T", "sv_return_wa_60T",
        "blank col1", "Woning_type", "bouwjaar_bins", "BJ_Cat", "Gebruiksoppervlak_bin", "indicatief_gas [m3]",
        "Total installed power, Q_a (kW)", "dimless heat power, Q_d/Q_a [-]",
    ]
    for c, naam in enumerate(kop, start=1):
        sheet1.cell(row=6, column=c, value=naam)
    for r, home in enumerate(homes, start=7):
        waarden = [
            r - 7, home, 0.2, 0.25, 0.01, 0.02, 5000.0, 1e7, 0.001, 2.0, 5e7, 0.005,
            55.0, 40.0, None, "tussenwoning", "1974-1982", "1974-1991", "80-89", 1200.0, 10.0, 0.5,
        ]
        for c, waarde in enumerate(waarden, start=1):
            sheet1.cell(row=r, column=c, value=waarde)
    wb.create_sheet("Sheet2")

    xlsx_buf = io.BytesIO()
    wb.save(xlsx_buf)
    xlsx_bytes = xlsx_buf.getvalue()

    with zipfile.ZipFile(pad, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr("boiler-temp.csv", boiler_inhoud)
        z.writestr("thermostat-temp.csv", thermo_inhoud)
        z.writestr("boiler-24h-stats.csv", etmaal_inhoud)
        z.writestr("README.md", readme_inhoud)
        z.writestr("resultaat-tabel-WarmingUP_KITE openbare versie.xlsx", xlsx_bytes)

    return pad


@pytest.fixture
def test_zip_gesorteerd(tmp_path: Path) -> Path:
    return bouw_test_zip(tmp_path / "test-data.zip", ["WU0", "WU1", "WU2"], rijen_per_woning=15)


@pytest.fixture
def test_zip_ongesorteerd(tmp_path: Path) -> Path:
    return bouw_test_zip(tmp_path / "test-data-ongesorteerd.zip", ["WU0", "WU1"], rijen_per_woning=12, ongesorteerd=True)
